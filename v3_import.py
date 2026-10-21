"""
Enrichisseur v3 — import d'un Excel de lancement dans la mémoire Notion (sans relancer la recherche)

Pour les lancements faits avant la mémoire des sociétés : l'Excel produit par l'outil (onglets Résultats et Dirigeants)
est retraduit en fiches mémoire. Les sociétés déjà en mémoire (même SIREN, domaine ou nom) sont laissées telles quelles.
Ensuite, le rattrapage ajoute le contrôle des contacts CRM.

Aperçu d'abord (rien n'est écrit), import réel ensuite.
"""
import io, re

from fastapi import APIRouter, File, Form, Request, UploadFile
from fastapi.responses import JSONResponse
from openpyxl import load_workbook

import v3_memoire as MEM
import v3_validation as VALID

router = APIRouter()

# en-têtes de l'onglet « Résultats » (voir excel() dans v3.py) → clés internes
COLONNES = {
    "ID Pipedrive": "id", "Société": "nom", "Déjà au CRM": "dans_crm", "Traitement": "mode", "Domaine": "domaine",
    "Site trouvé via": "source_site", "Contrôle du nom": "controle_nom", "Site joignable": "site_joignable",
    "SIREN CRM": "siren_crm", "État SIREN CRM": "etat_crm", "SIREN site": "siren_site", "État SIREN site": "etat_site",
    "Page où le SIREN a été trouvé": "page_siren", "SIREN par le nom (SIRENE)": "siren_nom", "Raison sociale INSEE": "raison_sociale",
    "Date fermeture": "date_fermeture", "Siège légal INSEE": "siege_insee", "Code NAF": "naf", "Activité (NAF)": "activite", "Secteur": "secteur",
    "LinkedIn (site)": "linkedin_site", "Statut proposé": "statut", "Action": "action", "Confiance": "confiance", "Note": "note_site",
    "Adresse CRM": "adresse_crm", "Adresse site": "adresse_site", "Page de l'adresse site": "page_adresse_site",
    "Page LinkedIn retenue": "linkedin", "Trouvée via": "source_linkedin", "Contrôle LinkedIn": "controle_linkedin",
    "Adresse LinkedIn": "adresse_linkedin", "Fiche Google": "fiche_google", "Adresse Google": "adresse_google",
    "Adresse proposée": "adresse_proposee", "Sources concordantes": "sources_adresse", "Action adresse": "action_adresse",
    "Confiance adresse": "confiance_adresse", "Verdict domiciliation": "verdict_domiciliation", "Pourquoi": "raison_domiciliation",
    "Effectif LinkedIn": "effectif_linkedin", "Évolution effectif LinkedIn": "evolution_effectif",
    "Croissance effectif (%)": "croissance_effectif", "Tranche déclarée": "tranche_linkedin", "Statut activité": "statut_activite"}
PERIMETRES = [("tout", "tout"), ("sociétés + dirigeants", "dirigeants"), ("dirigeants", "dirigeants"),
              ("sociétés", "societes"), ("societes", "societes"), ("siren", "siren")]


def _txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def lire_excel(nom_fichier: str, data: bytes) -> dict:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if "Résultats" not in wb.sheetnames:
        raise ValueError("ce fichier n'est pas un Excel de l'enrichisseur v3 (onglet « Résultats » absent)")
    titre = _txt(next(wb["Lisez-moi"].iter_rows(values_only=True))[0]) if "Lisez-moi" in wb.sheetnames else ""
    perim = next((p for mot, p in PERIMETRES if mot in titre.lower()), "")
    m = re.match(r"enrichisseur_v3_(\d{4}-\d{2}-\d{2})_(.+?)(_contacts)?\.xlsx$", nom_fichier.rsplit("/", 1)[-1])
    lancement = (m.group(2) if m else nom_fichier) + " (import Excel)"
    lignes = list(wb["Résultats"].iter_rows(values_only=True))
    entetes = [_txt(h) for h in lignes[0]]
    societes = []
    for x in lignes[1:]:
        r = {COLONNES[h]: _txt(v) for h, v in zip(entetes, x) if h in COLONNES}
        if not r.get("nom"):
            continue
        try:
            r["effectif_linkedin"] = int(float(r["effectif_linkedin"])) if r.get("effectif_linkedin") else ""
        except ValueError:
            r["effectif_linkedin"] = ""
        r["dirigeants"], r["source"] = [], f"import Excel ({lancement.replace(' (import Excel)', '')})"
        societes.append(r)
    if "Dirigeants" in wb.sheetnames:
        d_l = list(wb["Dirigeants"].iter_rows(values_only=True))
        h = [_txt(c) for c in d_l[0]]
        par_cle = {(s["id"] or s["nom"]): s for s in societes}
        for x in d_l[1:]:
            v = {k: _txt(c) for k, c in zip(h, x)}
            s = par_cle.get(v.get("ID Pipedrive") or v.get("Société"))
            if not s:
                continue
            s["dirigeants"].append({
                "prenom": v.get("Prénom", ""), "nom": v.get("Nom", ""), "nom_complet": f"{v.get('Prénom', '')} {v.get('Nom', '')}".strip(),
                "fonction": v.get("Fonction", ""), "role": v.get("Rôle", ""), "lieu": v.get("Lieu", ""), "idf": v.get("Île-de-France") == "oui",
                "linkedin": v.get("LinkedIn", ""), "email": v.get("E-mail", ""), "source_email": v.get("Source e-mail", ""),
                "dans_pipedrive": v.get("Déjà dans Pipedrive", "")})
    return {"lancement": lancement, "perimetre": perim, "societes": societes}


async def _trier(lu: dict) -> tuple:
    from v3 import norm_nom
    idx = await MEM.index(lambda *_: None, tout=True)
    neuves, deja = [], []
    for s in lu["societes"]:
        siren = s.get("siren_site") or s.get("siren_crm") or s.get("siren_nom") or ""
        (deja if any(MEM._meme(e, siren, s.get("domaine", ""), norm_nom(s["nom"])) for e in idx) else neuves).append(s)
    return neuves, deja


@router.post("/v3/validation/api/import")
async def importer(req: Request, fichier: UploadFile = File(...), mode: str = Form("apercu")):
    if (x := VALID._refus(req, "admin")):
        return x
    try:
        lu = lire_excel(fichier.filename or "", await fichier.read())
    except Exception as e:
        return JSONResponse({"erreur": f"lecture impossible : {e}"}, 400)
    if not lu["perimetre"]:
        return JSONResponse({"erreur": "choix du lancement illisible dans l'onglet Lisez-moi"}, 400)
    neuves, deja = await _trier(lu)
    rep = {"lancement": lu["lancement"], "perimetre": lu["perimetre"], "a_importer": [s["nom"] for s in neuves],
           "deja_en_memoire": [s["nom"] for s in deja], "dirigeants": sum(len(s["dirigeants"]) for s in neuves)}
    if mode != "importer" or not neuves:
        return rep | {"importe": 0}
    from v3 import norm_nom
    journal = []
    job = {"perimetre": lu["perimetre"], "nom": lu["lancement"], "etat": "fini", "resultats": neuves}
    await MEM.enregistrer(job, norm_nom, journal.append)
    return rep | {"importe": len(neuves), "journal": journal}
