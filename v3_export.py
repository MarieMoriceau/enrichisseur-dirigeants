"""
Enrichisseur v3 — export des validations au format d'import HubSpot (en attendant la migration de fin octobre)

Un seul Excel, cumulatif (tout ce qui est validé à ce jour) : HubSpot reconnaît les sociétés par leur domaine
et les contacts par leur e-mail, donc réimporter le même fichier met à jour au lieu de créer des doublons.
  · Sociétés          : sociétés validées (adresse retenue, effectif LinkedIn, date de validation société)
  · Contacts à créer   : nouveaux dirigeants validés par les négos (saisies manuelles comprises, date de validation contact)
  · Contacts existants : dirigeants déjà dans Pipedrive chez une autre société (décision admin : changer de société / deux sociétés)
  · Contacts CRM       : décisions de l'admin sur les contacts déjà au CRM (garder / marquer parti / à supprimer / à revoir, corrections)
  · Sociétés écartées
Envoyé par mail à COPIE_FICHIER_COMPLET (Marie) ou téléchargé depuis l'écran admin.
"""
import asyncio, io, re
from datetime import datetime, timezone

from fastapi import APIRouter, Request
from fastapi.responses import JSONResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font

import v3_memoire as MEM
import v3_suivi as SUIVI
import v3_validation as VALID

router = APIRouter()
_SEM = asyncio.Semaphore(4)


def _date(x: str) -> str:
    return (x or "")[:10]


def _prenom_nom(nom_complet: str, prenom: str = "", nom: str = "") -> tuple:
    if prenom or nom:
        return prenom, nom
    m = (nom_complet or "").strip().split(" ", 1)
    return (m[0], m[1]) if len(m) == 2 else ("", nom_complet or "")


async def _donnees() -> list:
    """Sociétés de la mémoire avec leur JSON complet (seulement celles qui ont une validation ou des contacts traités)."""
    toutes = await VALID._toutes()
    utiles = [s for s in toutes if s["validation"] != "À valider" or s["nouveaux"] or s["decisions"]]

    async def lire(s):
        async with _SEM:
            s["data"] = await MEM.lire(s["page"])
        VALID.appliquer_correction_siren(s)                 # SIREN corrigé à la main : la bonne fiche INSEE
        s["decisions"] = VALID._json(s["decisions"], [])
        s["nouveaux"] = VALID._json(s["nouveaux"], {})
        return s
    return await asyncio.gather(*(lire(s) for s in utiles))


def construire(socs: list) -> tuple:
    from v3 import decouper_adresse, feuille
    auj = datetime.now(timezone.utc).date().isoformat()
    validees = [s for s in socs if s["validation"] in ("Société validée", "Contacts validés")]
    ecartees = [s for s in socs if s["validation"] == "Écartée"]
    domaine_de = {s["page"]: s["domaine"] or (s["data"] or {}).get("domaine", "") for s in socs}

    # ── sociétés
    l_soc = []
    for s in validees:
        d = s["data"] or {}
        rue, cp, ville = decouper_adresse(s["adresse_retenue"])
        eff = d.get("effectif_linkedin", "")
        l_soc.append([domaine_de[s["page"]], s["nom"], s["siren"], d.get("linkedin") or d.get("linkedin_site", ""),
                      rue, cp, ville, s["source_retenue"], eff, _date(s["traitee"]) if eff != "" else "", d.get("croissance_effectif", ""),
                      d.get("statut_activite") or d.get("statut", ""), d.get("raison_sociale", ""), d.get("naf", ""), VALID.activite_de(d), d.get("secteur", ""),
                      re.sub(r"^\W+", "", d.get("verdict_domiciliation", "")).strip(), s.get("base_quali") or d.get("base_quali_crm", ""),
                      s.get("debut_bail", ""), "oui" if s.get("flexoffice") else "",
                      (s.get("operateur_flex") or "").replace("(aucun)", "").replace(" (probable)" if s.get("flexoffice") else "(probable)",
                                                              "" if s.get("flexoffice") else "(à confirmer)").strip(),
                      d.get("id", ""), _date(s["validee"])])

    # ── contacts (nouveaux validés par les négos)
    l_new, l_exist = [], []
    for s in socs:
        if s["validation"] == "Écartée":
            continue
        dirs = {VALID.cle_contact(x): x for x in VALID.nouveaux_dirigeants(s["data"] or {})}
        dec_a = {str(x.get("id")): x for x in s["decisions"] if x.get("type") == "ailleurs"}
        for cle, e in s["nouveaux"].items():
            if e.get("etat") != "prêt pour HubSpot":
                continue
            d = dirs.get(cle, {})
            f = VALID.fiche_nouveau(d, e) if d else {k: e.get("final_" + k, e.get(k, "")) for k in VALID.CHAMPS_MODIFIABLES}
            prenom, nom = _prenom_nom(e.get("nom_complet") or d.get("nom_complet", ""), e.get("prenom") or d.get("prenom", ""),
                                      e.get("nom") or d.get("nom", ""))
            commun = [prenom, nom, f.get("email", ""), f.get("telephone", ""), f.get("fonction", ""), f.get("role", ""),
                      f.get("linkedin", ""), domaine_de[s["page"]], s["nom"]]
            societe_ok = "oui" if s["validation"] in ("Société validée", "Contacts validés") else "non (adresse pas encore validée)"
            saisies = ", ".join(e.get("modifie", []))
            if e.get("action") == "existant":
                a = dec_a.get(str(e.get("fiche_existante")), {})
                l_exist.append(commun + [e.get("fiche_existante", ""), (d.get("ailleurs_crm") or {}).get("org_nom", "") or a.get("societe_actuelle", ""),
                                         a.get("action", "à décider"), _date(e.get("le")), societe_ok, saisies])
            else:
                l_new.append(commun + [d.get("id") or (s["data"] or {}).get("id", ""), _date(e.get("le")), societe_ok, saisies,
                                       e.get("modifie_par", "")])

    # ── contacts déjà au CRM (décisions admin)
    l_crm = []
    for s in socs:
        crm = {str(p.get("id")): p for p in (s["data"] or {}).get("contacts_crm", [])}
        for x in s["decisions"]:
            if (x.get("type") or "crm") != "crm":
                continue
            p = crm.get(str(x.get("id")), {})
            prenom, nom = _prenom_nom(x.get("nom_complet") or p.get("nom_complet", ""))
            l_crm.append([x.get("action", ""), prenom, nom, x.get("email") or p.get("email", ""), x.get("telephone") or p.get("telephone", ""),
                          x.get("fonction") or p.get("fonction", ""), x.get("role") or p.get("role", ""), x.get("linkedin") or p.get("linkedin", ""),
                          domaine_de[s["page"]], s["nom"], x.get("id", ""), x.get("statut") or p.get("statut", ""), p.get("detail", ""),
                          _date(x.get("le") or s["validee"])])
    ordre = {"à supprimer": 0, "marquer parti": 1, "à revoir": 2, "garder": 3}
    l_crm.sort(key=lambda l: (ordre.get(l[0], 3), l[9]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Lisez-moi"
    lignes = [f"ENRICHISSEUR V3 · export des validations pour HubSpot · {auj}",
              "Fichier cumulatif : tout ce qui est validé à ce jour. HubSpot reconnaît les sociétés par le domaine et les contacts par l'e-mail :",
              "réimporter le même fichier met à jour les fiches au lieu de créer des doublons.", "",
              f"Sociétés validées : {len(l_soc)}",
              f"Contacts à créer (validés par les négos) : {len(l_new)}",
              f"Contacts déjà dans Pipedrive ailleurs (à rattacher ou déplacer) : {len(l_exist)}",
              f"Décisions sur les contacts CRM : {len(l_crm)} (dont {sum(1 for l in l_crm if l[0] == 'marquer parti')} à marquer partis, {sum(1 for l in l_crm if l[0] == 'à supprimer')} à supprimer)",
              f"Sociétés écartées : {len(ecartees)}, dont {sum(1 for s in ecartees if s.get('decision_ecart') == 'À supprimer de la base')} à supprimer de Pipedrive (onglet dédié)", "",
              "Contacts partis : contacts retrouvés dans une nouvelle société et validés. Mise à jour de la fiche existante :",
              "  importer avec « Email actuel » comme clé ; le nouvel e-mail va dans une propriété (« nouvel_email ») ou se corrige à la main.",
              "  Les nouvelles sociétés sont dans « Sociétés des partis » (à importer avant, pour l'association par domaine).", "",
              "Ordre d'import conseillé dans HubSpot : 1. Sociétés  2. Contacts à créer (associés par Company Domain Name)",
              "  3. Contacts existants et Contacts CRM : à traiter à la main ou en import de mise à jour (clé = e-mail).",
              "Propriétés HubSpot à créer avant l'import : date_validation_societe, date_validation_contact, effectif_linkedin,",
              "  statut_activite, domiciliation, role_dirigeant. Les colonnes « Date validation » sont prévues pour elles.",
              "Les contacts dont la société n'est pas encore validée sont signalés (colonne « Société validée ? »)."]
    for l in lignes:
        ws.append([l])
    ws["A1"].font = Font(name="Arial", bold=True, size=13, color="4A1E5C")
    ws.column_dimensions["A"].width = 120
    feuille(wb, "Sociétés", ["Company Domain Name", "Name", "SIREN", "LinkedIn Company Page", "Street Address", "Postal Code", "City",
                             "Source de l'adresse", "Effectif LinkedIn", "Date relevé effectif", "Croissance effectif (%)", "STATUT ACTIVITÉ",
                             "Raison sociale", "CODE NAF", "ACTIVITÉ", "SECTEUR", "DOMICILIATION", "Base quali", "Début de bail", "Flexoffice", "Opérateur flex", "ID Pipedrive", "Date validation société"], l_soc)
    feuille(wb, "Contacts à créer", ["First Name", "Last Name", "Email", "Phone Number", "Job Title", "Rôle", "LinkedIn URL",
                                     "Company Domain Name", "Société", "ID Pipedrive société", "Date validation contact", "Société validée ?",
                                     "Champs saisis à la main", "Saisi par"], l_new)
    feuille(wb, "Contacts existants", ["First Name", "Last Name", "Email", "Phone Number", "Job Title", "Rôle", "LinkedIn URL",
                                       "Company Domain Name", "Nouvelle société", "ID personne Pipedrive", "Fiché aujourd'hui chez",
                                       "Décision admin", "Date validation contact", "Société validée ?", "Champs saisis à la main"], l_exist)
    feuille(wb, "Contacts CRM", ["Décision", "First Name", "Last Name", "Email", "Phone Number", "Job Title", "Rôle", "LinkedIn URL",
                                 "Company Domain Name", "Société", "ID personne Pipedrive", "Contrôle", "Détail du contrôle", "Date validation"], l_crm)
    # ── contacts partis retrouvés et validés : mise à jour de la fiche existante (nouvelle société, poste, e-mail)
    import v3_partis as PARTIS
    partis = [x for x in PARTIS.liste(socs) if x.get("etat") == "validé"]
    l_partis, l_soc_partis, vus = [], [], set()
    for x in partis:
        prenom, nom = _prenom_nom(x.get("nom_complet", ""))
        l_partis.append([x.get("ancien_email", ""), x.get("email", ""), prenom, nom, x.get("fonction", ""), x.get("role", ""),
                         x.get("linkedin", ""), x.get("domaine", ""), x.get("nouvelle_societe", ""), x.get("ancienne_societe", ""),
                         x.get("id", ""), " ; ".join(x.get("alertes", [])), _date(x.get("decide_le") or x.get("le"))])
        cle = x.get("domaine") or x.get("nouvelle_societe")
        if cle and cle not in vus:
            vus.add(cle)
            rue, cp, ville = decouper_adresse(x.get("adresse", ""))
            l_soc_partis.append([x.get("domaine", ""), x.get("nouvelle_societe", ""), x.get("linkedin_societe", ""),
                                 rue, cp, ville, x.get("effectif", "")])
    feuille(wb, "Contacts partis", ["Email actuel (clé HubSpot)", "Nouvel e-mail", "First Name", "Last Name", "Job Title", "Rôle",
                                    "LinkedIn URL", "Company Domain Name", "Nouvelle société", "Ancienne société",
                                    "ID personne Pipedrive", "Alertes", "Date validation"], l_partis)
    feuille(wb, "Sociétés des partis", ["Company Domain Name", "Name", "LinkedIn Company Page", "Street Address", "Postal Code",
                                        "City", "Effectif LinkedIn"], l_soc_partis)
    a_suppr = [s for s in ecartees if s.get("decision_ecart") == "À supprimer de la base"]
    feuille(wb, "Sociétés écartées", ["Société", "Domaine", "SIREN", "ID Pipedrive", "Décision", "Motif", "Écartée le"],
            [[s["nom"], domaine_de[s["page"]], s["siren"], (s["data"] or {}).get("id", ""), s.get("decision_ecart") or "Hors cible HubSpot",
              s.get("motif_ecart", ""), _date(s["validee"])] for s in ecartees])
    feuille(wb, "À supprimer (Pipedrive, HubSpot)", ["Société", "ID Pipedrive", "SIREN", "Domaine", "Motif", "Décidé le"],
            [[s["nom"], (s["data"] or {}).get("id", ""), s["siren"], domaine_de[s["page"]], s.get("motif_ecart", ""), _date(s["validee"])]
             for s in a_suppr])
    b = io.BytesIO()
    wb.save(b)
    resume = {"societes": len(l_soc), "contacts": len(l_new), "existants": len(l_exist), "crm": len(l_crm), "ecartees": len(ecartees),
              "a_supprimer": len(a_suppr), "partis": len(l_partis)}
    return b.getvalue(), resume


def nom_fichier() -> str:
    return f"enrichisseur_v3_export_hubspot_{datetime.now(timezone.utc):%Y-%m-%d}.xlsx"


@router.get("/v3/validation/api/export")
async def telecharger(req: Request):
    if (x := VALID._refus(req, "admin")):
        return x
    xl, _ = construire(await _donnees())
    return StreamingResponse(io.BytesIO(xl), headers={"Content-Disposition": f'attachment; filename="{nom_fichier()}"'},
                             media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@router.post("/v3/validation/api/export/mail")
async def envoyer(req: Request):
    if (x := VALID._refus(req, "admin")):
        return x
    if not (SUIVI.SMTP_USER and SUIVI.SMTP_PASS and SUIVI.COPIE_COMPLETE):
        return JSONResponse({"erreur": "envoi de mail non configuré (SMTP_USER / SMTP_PASS / COPIE_FICHIER_COMPLET)"}, 500)
    xl, r = construire(await _donnees())
    corps = "\n".join(["Bonjour,", "", "Voici l'export des validations de l'enrichisseur v3, prêt pour l'import HubSpot.", "",
                       f"Sociétés validées : {r['societes']}", f"Contacts à créer : {r['contacts']}",
                       f"Contacts déjà dans Pipedrive ailleurs : {r['existants']}", f"Décisions sur les contacts CRM : {r['crm']}",
                       f"Sociétés écartées : {r['ecartees']} (dont {r.get('a_supprimer', 0)} à supprimer de Pipedrive)",
                       f"Contacts partis retrouvés (mise à jour) : {r.get('partis', 0)}", "",
                       "Le fichier est cumulatif : tout ce qui est validé à ce jour. Le mode d'emploi est dans l'onglet Lisez-moi.", "",
                       "Enrichisseur v3"])
    try:
        await asyncio.to_thread(SUIVI._envoyer, SUIVI.COPIE_COMPLETE, f"Enrichisseur v3 · export HubSpot du {datetime.now():%d/%m/%Y}",
                                corps, xl, nom_fichier())
    except Exception as e:
        return JSONResponse({"erreur": f"envoi impossible : {e}"}, 502)
    return r | {"envoye_a": SUIVI.COPIE_COMPLETE}


# ───────────────────────── liste de phoning d'un fichier lancé ─────────────────────────
_SUITE = re.compile(r"( \((suite \d+/\d+|relance[^)]*)\))+$")


@router.get("/v3/contacts/api/phoning")
async def phoning(req: Request, lancement: str = ""):
    """Excel de phoning d'un fichier lancé (toutes ses suites) : nouveaux contacts validés (« prêt pour HubSpot »)
    + contacts déjà dans Pipedrive (sauf ceux marqués partis / à supprimer).
    Admin : tout ; négo : ses sociétés seulement (même règle que son écran)."""
    if (x := VALID._refus(req, "negos")):
        return x
    base = _SUITE.sub("", lancement.strip())
    if len(base) < 3:
        return JSONResponse({"erreur": "précisez le fichier (?lancement=…)"}, 400)
    qui = await VALID.moi(req)
    if not qui.get("role"):
        return JSONResponse({"erreur": "accès refusé : reconnectez-vous"}, 403)
    socs = [s for s in await VALID._toutes()
            if any(_SUITE.sub("", l.strip()) == base for l in (s.get("lancement") or "").split(";"))
            and s["validation"] != "Écartée" and VALID.visible_par(qui, s)]

    async def lire(s):
        for essai in range(3):
            try:
                async with _SEM:
                    s["data"] = await MEM.lire(s["page"])
                break
            except Exception:
                s["data"] = {}
                await asyncio.sleep(2 + 2 * essai)
        s["nouveaux"] = VALID._json(s["nouveaux"], {})
        s["decisions"] = VALID._json(s["decisions"], [])
        return s
    socs = await asyncio.gather(*(lire(s) for s in socs))
    lignes = []
    for s in sorted(socs, key=lambda s: (s["nom"] or "").lower()):
        d0 = s["data"] or {}
        adresse = s.get("adresse_retenue") or d0.get("adresse_proposee") or d0.get("adresse_crm") or ""
        soc = [s["nom"], s.get("domaine") or d0.get("domaine", ""), adresse, d0.get("effectif_linkedin", "")]
        # nouveaux contacts validés
        dirs = {VALID.cle_contact(x): x for x in VALID.nouveaux_dirigeants(d0)}
        for cle, e in s["nouveaux"].items():
            if e.get("etat") != "prêt pour HubSpot":
                continue
            d = dirs.get(cle, {})
            f = VALID.fiche_nouveau(d, e) if d else {k: e.get("final_" + k, e.get(k, "")) for k in VALID.CHAMPS_MODIFIABLES}
            prenom, nom = _prenom_nom(e.get("nom_complet") or d.get("nom_complet", ""), e.get("prenom") or d.get("prenom", ""),
                                      e.get("nom") or d.get("nom", ""))
            autre = (d.get("ailleurs_crm") or {}).get("org_nom", "")
            lignes.append(soc + ["Nouveau contact", prenom, nom, f.get("fonction", ""), f.get("role", ""), f.get("email", ""),
                                 f.get("telephone", ""), f.get("linkedin", ""),
                                 f"validé le {_date(e.get('le'))}" + (f" · déjà dans Pipedrive chez {autre}" if autre else ""),
                                 s.get("negos", ""), "", ""])
        # contacts déjà dans Pipedrive (corrections de l'écran admin comprises)
        dec = {str(x.get("id")): x for x in s["decisions"] if (x.get("type") or "crm") == "crm"}
        for p in d0.get("contacts_crm", []):
            x = dec.get(str(p.get("id")), {})
            if x.get("action") in ("marquer parti", "à supprimer") or p.get("statut") == "probablement parti" and not x:
                continue
            v = lambda k: (x.get(k) or "").strip() or p.get(k, "")
            prenom, nom = _prenom_nom(p.get("nom_complet", ""))
            etat = x.get("action") or p.get("statut") or "non contrôlé"
            lignes.append(soc + ["Pipedrive", prenom, nom, v("fonction"), v("role"), v("email"), v("telephone"), v("linkedin"),
                                 f"{etat} · fiche Pipedrive {p.get('id', '')}" + (f" (doublon {p['fiche_doublon']})" if p.get("fiche_doublon") else ""),
                                 s.get("negos", ""), "", ""])
    from v3 import feuille
    wb = Workbook()
    wb.remove(wb.active)
    feuille(wb, "Phoning", ["Société", "Site", "Adresse des bureaux", "Effectif LinkedIn", "Source", "Prénom", "Nom", "Fonction",
                            "Rôle", "E-mail", "Téléphone", "LinkedIn", "Statut", "Négo", "Date d'appel", "Résultat / notes"], lignes)
    buf = io.BytesIO(); wb.save(buf)
    nom = re.sub(r"[^A-Za-z0-9_-]+", "_", base)[:50].strip("_") or "fichier"
    return StreamingResponse(io.BytesIO(buf.getvalue()), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="phoning_{nom}_{datetime.now():%Y-%m-%d}.xlsx"'})
