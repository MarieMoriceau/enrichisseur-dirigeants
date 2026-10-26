"""Liste de phoning d'un fichier lancé. Lancer : python3 -m pytest tests/"""
import io, json, os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import v3, v3_export as EXPORT, v3_validation as VALID, v3_memoire as MEM


def test_phoning(monkeypatch):
    nouveaux = {"jean dupont": {"etat": "prêt pour HubSpot", "nom_complet": "Jean Dupont", "prenom": "Jean", "nom": "Dupont",
                                "final_email": "jd@x.fr", "final_fonction": "CEO", "final_role": "1 CEO", "le": "2026-10-25T10:00:00"},
                "paul": {"etat": "écarté", "nom_complet": "Paul X"}}
    async def toutes():
        return [{"page": "p1", "nom": "Botify", "lancement": "locataires_flex_paris (suite 1/2) ; autre", "validation": "À valider",
                 "nouveaux": json.dumps(nouveaux, ensure_ascii=False), "negos": "Richard", "adresse_retenue": ""},
                {"page": "p2", "nom": "Hors fichier", "lancement": "autre", "validation": "À valider",
                 "nouveaux": json.dumps(nouveaux, ensure_ascii=False)}]
    async def lire(p):
        return {"domaine": "botify.com", "adresse_proposee": "7 rue de Madrid, 75008 Paris", "effectif_linkedin": 274, "dirigeants": []}
    async def moi(req):
        return {"role": "admin", "mail": "m@x.fr", "membre": {}, "equipe": []}
    monkeypatch.setattr(VALID, "_toutes", toutes); monkeypatch.setattr(MEM, "lire", lire); monkeypatch.setattr(VALID, "moi", moi)
    monkeypatch.setattr(VALID, "_refus", lambda req, besoin: None)
    app = FastAPI(); app.include_router(EXPORT.router)
    r = TestClient(app).get("/v3/contacts/api/phoning", params={"lancement": "locataires_flex_paris (suite 1/2)"})
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 2 and rows[1][0] == "Botify" and rows[1][4] == "Jean" and rows[1][8] == "jd@x.fr"
    assert rows[1][2].startswith("7 rue de Madrid")
